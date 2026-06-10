"""
Aplicación GUI para consultar el Tipo de Cambio Oficial de SUNAT.

Modos de consulta:
  - Hoy
  - Rango de días (desde / hasta)
  - Mes completo (mes / año)

Funciones:
  - Copiar filas seleccionadas o todas al portapapeles
  - Exportar a archivo Excel (.xlsx) con formato
  - Caché local SQLite (consultas repetidas son instantáneas)
  - Filtros rápidos pre-armados (Hoy, Ayer, Últ. 7d, etc.)
  - Recuerda la última consulta y posición de ventana

Fuente de datos: https://api.apis.net.pe/v1/tipo-cambio-sunat
(API pública que replica el tipo de cambio oficial publicado por SUNAT)

La capa de datos vive en core.py; este archivo sólo contiene el GUI.
"""

from __future__ import annotations

import calendar
import queue
import sys
import threading
import tkinter as tk
from datetime import date, datetime, timedelta
from tkinter import filedialog, messagebox, ttk
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import (
    CACHE_DB,
    MIN_YEAR,
    MAX_YEAR,
    cache_clear,
    cache_count,
    cache_date_range,
    fetch_day,
    fetch_month,
    fetch_range,
    load_state,
    save_state,
)


# ---------------------------------------------------------------------------
# Interfaz gráfica
# ---------------------------------------------------------------------------

class DateEntry(ttk.Frame):
    """
    Selector de fecha con tres Spinbox (DD / MM / YYYY).

    - Sin dependencias externas.
    - Los Spinbox limitan el rango con flechas; el día se ajusta solo cuando
      cambia el mes o el año (p. ej. 31 -> 28 en febrero).
    - set_date(d) y get_date() permiten programar/recuperar el valor.
    """

    def __init__(
        self,
        parent: tk.Misc,
        default: date | None = None,
        min_year: int = MIN_YEAR,
        max_year: int | None = None,
        **kwargs,
    ) -> None:
        super().__init__(parent, **kwargs)
        if max_year is None:
            max_year = date.today().year
        if default is None:
            default = date.today()
        self._min_year = min_year
        self._max_year = max_year

        self._day_var = tk.StringVar(value=f"{default.day:02d}")
        self._month_var = tk.StringVar(value=f"{default.month:02d}")
        self._year_var = tk.StringVar(value=str(default.year))

        self._day = ttk.Spinbox(
            self,
            from_=1,
            to=31,
            width=3,
            justify="center",
            textvariable=self._day_var,
        )
        self._day.pack(side=tk.LEFT)
        ttk.Label(self, text="/").pack(side=tk.LEFT, padx=2)
        self._month = ttk.Spinbox(
            self,
            from_=1,
            to=12,
            width=3,
            justify="center",
            textvariable=self._month_var,
        )
        self._month.pack(side=tk.LEFT)
        ttk.Label(self, text="/").pack(side=tk.LEFT, padx=2)
        self._year = ttk.Spinbox(
            self,
            from_=min_year,
            to=max_year,
            width=5,
            justify="center",
            textvariable=self._year_var,
        )
        self._year.pack(side=tk.LEFT)

        # Ajustar el día máximo cuando cambia mes/año
        self._month_var.trace_add("write", lambda *_: self._adjust_day())
        self._year_var.trace_add("write", lambda *_: self._adjust_day())
        # Al perder foco, normalizar lo escrito
        for w in (self._day, self._month, self._year):
            w.bind("<FocusOut>", lambda _e: self._normalize())
        self._adjust_day()

    # -- lógica interna -------------------------------------------------------
    def _adjust_day(self) -> None:
        """Recalcula el día máximo permitido según mes/año actuales."""
        try:
            m = int(self._month_var.get())
            y = int(self._year_var.get())
        except (ValueError, tk.TclError):
            return
        try:
            last_day = calendar.monthrange(y, m)[1]
        except ValueError:
            return
        self._day.configure(to=last_day)
        try:
            cur = int(self._day_var.get())
            if cur > last_day:
                self._day_var.set(f"{last_day:02d}")
        except (ValueError, tk.TclError):
            pass

    def _normalize(self) -> None:
        """Rellena con ceros a la izquierda si el usuario escribió 1 dígito."""
        for var, width in (
            (self._day_var, 2),
            (self._month_var, 2),
        ):
            v = var.get().strip()
            if v.isdigit() and len(v) < width:
                var.set(v.zfill(width))
        try:
            y = int(self._year_var.get())
            if not (self._min_year <= y <= self._max_year):
                self._year_var.set(str(min(max(y, self._min_year), self._max_year)))
        except (ValueError, tk.TclError):
            pass

    # -- API pública ----------------------------------------------------------
    def get_date(self) -> date | None:
        try:
            d = int(self._day_var.get())
            m = int(self._month_var.get())
            y = int(self._year_var.get())
            return date(y, m, d)
        except (ValueError, tk.TclError):
            return None

    def set_date(self, value: date) -> None:
        # Importante: año y mes antes que el día, para que _adjust_day
        # (disparado por los trace_add) calcule el día máximo correcto.
        # Si fuera al revés, el día se clamparía con el mes/año anterior.
        self._year_var.set(str(value.year))
        self._month_var.set(f"{value.month:02d}")
        self._adjust_day()
        self._day_var.set(f"{value.day:02d}")
        self._adjust_day()


class SunatTCApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Tipo de Cambio SUNAT - Jose Miguel Maldonado Garcia")

        # Cargar estado persistido (ventana, modo, fechas)
        self.state = load_state()
        self.root.geometry(self.state.get("geometry", "850x600"))
        self.root.minsize(850, 600)
        self.root.protocol("WM_DELETE_WINDOW", self._on_close)

        self.data: list[dict] = []
        self.worker: threading.Thread | None = None
        self._cancelled = False
        self.progress_queue: queue.Queue[str] = queue.Queue()

        self._configure_style()
        self._build_ui()
        self._restore_state()
        self._on_mode_change()
        self._update_cache_label()

    # -- estilo ---------------------------------------------------------------
    def _configure_style(self) -> None:
        style = ttk.Style()
        try:
            style.theme_use("vista")  # nativo en Windows
        except tk.TclError:
            pass
        style.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
        style.configure("Treeview", font=("Segoe UI", 10), rowheight=24)
        style.configure("TButton", padding=(10, 5))
        style.configure("Accent.TButton", font=("Segoe UI", 10, "bold"))

    # -- construcción UI ------------------------------------------------------
    def _build_ui(self) -> None:
        # Frame superior: modo + fechas + acciones.
        # Usamos una sola columna (column=0) con sticky distinto por fila para
        # que los radio buttons NO se desplacen cuando cambian los controles
        # de fecha de abajo.
        top = ttk.LabelFrame(self.root, text="Consulta", padding=12)
        top.pack(side=tk.TOP, fill=tk.X, padx=12, pady=(12, 6))
        top.columnconfigure(0, weight=1)

        # --- Modo (row 0) ---
        self.mode = tk.StringVar(value="today")
        mode_box = ttk.Frame(top)
        mode_box.grid(row=0, column=0, sticky="w", pady=(0, 6))
        modes = [("Hoy", "today"), ("Rango de días", "range"), ("Mes completo", "month")]
        for label, val in modes:
            ttk.Radiobutton(
                mode_box,
                text=label,
                value=val,
                variable=self.mode,
                command=self._on_mode_change,
            ).pack(side=tk.LEFT, padx=(0, 14))

        # --- Controles de fecha (row 1) ---
        self.date_box = ttk.Frame(top)
        self.date_box.grid(row=1, column=0, sticky="w", pady=4)

        # -- Hoy
        self.today_label = ttk.Label(
            self.date_box,
            text=f"Mostrará el tipo de cambio del {date.today().strftime('%d/%m/%Y')}",
        )

        # -- Rango
        self.range_frame = ttk.Frame(self.date_box)
        ttk.Label(self.range_frame, text="Desde:").grid(row=0, column=0, padx=(0, 4))
        self.start_date = DateEntry(
            self.range_frame,
            default=date.today().replace(day=1),
        )
        self.start_date.grid(row=0, column=1, padx=(0, 14))
        ttk.Label(self.range_frame, text="Hasta:").grid(row=0, column=2, padx=(0, 4))
        self.end_date = DateEntry(
            self.range_frame,
            default=date.today(),
        )
        self.end_date.grid(row=0, column=3, padx=(0, 4))

        # -- Mes completo
        self.month_frame = ttk.Frame(self.date_box)
        ttk.Label(self.month_frame, text="Mes:").grid(row=0, column=0, padx=(0, 4))
        self.month_spin = ttk.Spinbox(
            self.month_frame, from_=1, to=12, width=4, justify="center"
        )
        self.month_spin.grid(row=0, column=1, padx=(0, 14))
        self.month_spin.set(date.today().month)
        ttk.Label(self.month_frame, text="Año:").grid(row=0, column=2, padx=(0, 4))
        self.year_spin = ttk.Spinbox(
            self.month_frame,
            from_=MIN_YEAR,
            to=MAX_YEAR,
            width=6,
            justify="center",
        )
        self.year_spin.grid(row=0, column=3, padx=(0, 4))
        self.year_spin.set(date.today().year)

        # --- Filtros rápidos (row 2) ---
        quick = ttk.Frame(top)
        quick.grid(row=2, column=0, sticky="w", pady=(8, 0))
        ttk.Label(quick, text="Filtros rápidos:").pack(side=tk.LEFT, padx=(0, 6))
        for label, kind in [
            ("Hoy", "today"),
            ("Ayer", "yesterday"),
            ("Últ. 7 días", "last7"),
            ("Últ. 30 días", "last30"),
            ("Este mes", "this_month"),
            ("Mes pasado", "last_month"),
        ]:
            ttk.Button(
                quick,
                text=label,
                command=lambda k=kind: self._apply_quick_filter(k),
            ).pack(side=tk.LEFT, padx=(0, 4))

        # --- Acciones (row 3) ---
        actions = ttk.Frame(top)
        actions.grid(row=3, column=0, sticky="we", pady=(10, 0))
        self.fetch_btn = ttk.Button(
            actions, text="Consultar", style="Accent.TButton", command=self._on_fetch
        )
        self.fetch_btn.pack(side=tk.LEFT)
        self.progress = ttk.Progressbar(actions, mode="indeterminate", length=200)
        self.status = ttk.Label(actions, text="Listo", foreground="#555")
        self.status.pack(side=tk.LEFT, padx=12)

        # Tabla de resultados
        table_frame = ttk.LabelFrame(self.root, text="Resultados", padding=6)
        table_frame.pack(side=tk.TOP, fill=tk.BOTH, expand=True, padx=12, pady=6)

        cols = ("fecha", "compra", "venta")
        self.tree = ttk.Treeview(
            table_frame, columns=cols, show="headings", selectmode="extended"
        )
        self.tree.heading("fecha", text="Fecha", anchor="center")
        self.tree.heading("compra", text="Compra (S/)", anchor="center")
        self.tree.heading("venta", text="Venta (S/)", anchor="center")
        self.tree.column("fecha", width=160, anchor="center")
        self.tree.column("compra", width=140, anchor="e")
        self.tree.column("venta", width=140, anchor="e")

        # Tags con colores alternados
        self.tree.tag_configure("odd", background="#f7f7fa")
        self.tree.tag_configure("even", background="#ffffff")

        vsb = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="we")
        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)

        # Menú contextual (clic derecho)
        self.tree.bind("<Button-3>", self._show_context_menu)
        self.context_menu = tk.Menu(self.root, tearoff=0)
        self.context_menu.add_command(
            label="Copiar selección", command=self._copy_selection
        )
        self.context_menu.add_command(label="Copiar todo", command=self._copy_all)
        self.context_menu.add_separator()
        self.context_menu.add_command(
            label="Exportar a Excel…", command=self._export_excel
        )

        # Atajos de teclado
        self.tree.bind("<Control-c>", lambda e: self._copy_selection())
        self.tree.bind("<Control-C>", lambda e: self._copy_selection())
        self.root.bind_all("<Control-e>", lambda e: self._export_excel())
        self.root.bind_all("<Control-E>", lambda e: self._export_excel())

        # Barra inferior
        bottom = ttk.Frame(self.root, padding=(12, 4, 12, 12))
        bottom.pack(side=tk.BOTTOM, fill=tk.X)
        ttk.Button(bottom, text="Copiar selección", command=self._copy_selection).pack(
            side=tk.LEFT
        )
        ttk.Button(bottom, text="Copiar todo", command=self._copy_all).pack(
            side=tk.LEFT, padx=6
        )
        ttk.Button(
            bottom, text="Limpiar caché",
            command=self._on_clear_cache,
        ).pack(side=tk.LEFT, padx=(6, 0))
        ttk.Button(
            bottom, text="Exportar a Excel…", command=self._export_excel
        ).pack(side=tk.RIGHT)

        self.count_label = ttk.Label(bottom, text="", foreground="#555")
        self.count_label.pack(side=tk.RIGHT, padx=12)

        self.cache_label = ttk.Label(bottom, text="", foreground="#888", cursor="hand2")
        self.cache_label.pack(side=tk.RIGHT, padx=12)
        self.cache_label.bind("<Button-1>", lambda _e: self._show_cache_info())

    # -- manejo de modo -------------------------------------------------------
    def _on_mode_change(self) -> None:
        # Ocultar todos
        for w in (self.today_label, self.range_frame, self.month_frame):
            w.pack_forget()
        m = self.mode.get()
        if m == "today":
            self.today_label.pack(side=tk.LEFT)
        elif m == "range":
            self.range_frame.pack(side=tk.LEFT)
        else:
            self.month_frame.pack(side=tk.LEFT)

    # -- filtros rápidos ------------------------------------------------------
    def _apply_quick_filter(self, kind: str) -> None:
        """Aplica un preset: ajusta modo + fechas y dispara la consulta."""
        today = date.today()
        if kind == "today":
            self.mode.set("today")
            self._on_mode_change()
        elif kind == "yesterday":
            d = today - timedelta(days=1)
            self.mode.set("range")
            self._on_mode_change()
            self.start_date.set_date(d)
            self.end_date.set_date(d)
        elif kind == "last7":
            self.mode.set("range")
            self._on_mode_change()
            self.start_date.set_date(today - timedelta(days=6))
            self.end_date.set_date(today)
        elif kind == "last30":
            self.mode.set("range")
            self._on_mode_change()
            self.start_date.set_date(today - timedelta(days=29))
            self.end_date.set_date(today)
        elif kind == "this_month":
            self.mode.set("month")
            self._on_mode_change()
            self.month_spin.set(today.month)
            self.year_spin.set(today.year)
        elif kind == "last_month":
            first_this = today.replace(day=1)
            last_prev = first_this - timedelta(days=1)
            self.mode.set("month")
            self._on_mode_change()
            self.month_spin.set(last_prev.month)
            self.year_spin.set(last_prev.year)
        # Dispara la consulta
        self._on_fetch()

    # -- persistencia de estado -----------------------------------------------
    def _restore_state(self) -> None:
        """Restaura modo + fechas desde el state cargado."""
        mode = self.state.get("mode", "today")
        if mode not in ("today", "range", "month"):
            mode = "today"
        self.mode.set(mode)
        if mode == "range":
            rs = self.state.get("range_start")
            re_ = self.state.get("range_end")
            if rs:
                try:
                    self.start_date.set_date(date.fromisoformat(rs))
                except ValueError:
                    pass
            if re_:
                try:
                    self.end_date.set_date(date.fromisoformat(re_))
                except ValueError:
                    pass
        elif mode == "month":
            td = date.today()
            try:
                self.month_spin.set(int(self.state.get("month", td.month)))
                self.year_spin.set(int(self.state.get("year", td.year)))
            except (ValueError, TypeError):
                pass

    def _on_close(self) -> None:
        """Persiste geometría + modo + fechas antes de cerrar."""
        self.state["geometry"] = self.root.geometry()
        self.state["mode"] = self.mode.get()
        if self.mode.get() == "range":
            sd = self.start_date.get_date()
            ed = self.end_date.get_date()
            if sd:
                self.state["range_start"] = sd.isoformat()
            if ed:
                self.state["range_end"] = ed.isoformat()
        elif self.mode.get() == "month":
            try:
                self.state["month"] = int(self.month_spin.get())
                self.state["year"] = int(self.year_spin.get())
            except (ValueError, tk.TclError):
                pass
        save_state(self.state)
        self.root.destroy()

    # -- caché ----------------------------------------------------------------
    def _update_cache_label(self) -> None:
        n = cache_count()
        dmin, dmax = cache_date_range()
        if n == 0:
            self.cache_label.configure(text="Caché: vacío")
        else:
            self.cache_label.configure(
                text=f"Caché: {n} días ({dmin} → {dmax})"
            )

    def _on_clear_cache(self) -> None:
        n = cache_count()
        if n == 0:
            messagebox.showinfo("Caché", "El caché ya está vacío.")
            return
        if not messagebox.askyesno(
            "Limpiar caché",
            f"Vas a borrar {n} entradas del caché.\n"
            "Las próximas consultas a fechas pasadas irán a la API.\n\n"
            "¿Continuar?",
        ):
            return
        cache_clear()
        self._update_cache_label()
        self.status.configure(text="Caché borrado.", foreground="#2da44e")

    def _show_cache_info(self) -> None:
        n = cache_count()
        dmin, dmax = cache_date_range()
        msg = (
            f"Entradas: {n}\n"
            f"Rango: {dmin or '-'}  →  {dmax or '-'}\n\n"
            f"Ubicación:\n{CACHE_DB}"
        )
        if messagebox.askyesno("Caché local", msg + "\n\n¿Abrir carpeta?"):
            try:
                import os, subprocess
                folder = str(CACHE_DB.parent)
                if hasattr(os, "startfile"):
                    os.startfile(folder)  # type: ignore[attr-defined]
                else:
                    subprocess.Popen(["xdg-open", folder])
            except Exception:
                pass

    def _show_context_menu(self, event) -> None:
        if not self.data:
            return
        try:
            self.context_menu.tk_popup(event.x_root, event.y_root)
        finally:
            self.context_menu.grab_release()

    # -- fetching -------------------------------------------------------------
    def _set_busy(self, busy: bool, message: str = "") -> None:
        if busy:
            self.fetch_btn.configure(state="disabled")
            self.progress.pack(side=tk.LEFT, padx=10)
            self.progress.start(12)
            self.status.configure(text=message, foreground="#1f6feb")
        else:
            self.fetch_btn.configure(state="normal")
            self.progress.stop()
            self.progress.pack_forget()
            self.status.configure(text=message, foreground="#2da44e")

    def _validate_inputs(self) -> tuple[bool, str, dict]:
        """Valida los inputs según el modo. Devuelve (ok, msg_error, params)."""
        m = self.mode.get()
        if m == "today":
            return True, "", {"kind": "today"}

        if m == "range":
            sd = self.start_date.get_date()
            ed = self.end_date.get_date()
            if sd is None or ed is None:
                return False, "Alguna de las fechas es inválida.", {}
            if sd > ed:
                return False, "'Desde' no puede ser mayor que 'Hasta'.", {}
            if ed > date.today():
                return False, "La fecha 'Hasta' no puede ser futura.", {}
            if (ed - sd).days > 730:
                return False, "El rango máximo es de 2 años (730 días).", {}
            return True, "", {"kind": "range", "start": sd, "end": ed}

        # month
        try:
            mo = int(self.month_spin.get())
            yr = int(self.year_spin.get())
        except ValueError:
            return False, "Mes y año deben ser numéricos.", {}
        if not (1 <= mo <= 12):
            return False, "El mes debe estar entre 1 y 12.", {}
        if not (MIN_YEAR <= yr <= MAX_YEAR):
            return False, f"El año debe estar entre {MIN_YEAR} y {MAX_YEAR}.", {}
        return True, "", {"kind": "month", "year": yr, "month": mo}

    def _on_fetch(self) -> None:
        if self.worker and self.worker.is_alive():
            return
        ok, err, params = self._validate_inputs()
        if not ok:
            messagebox.showwarning("Datos inválidos", err)
            return
        self._cancelled = False
        self._set_busy(True, "Consultando SUNAT…")
        self._clear_table()
        self.worker = threading.Thread(
            target=self._fetch_worker, args=(params, self._on_fetch_done), daemon=True
        )
        self.worker.start()
        self._poll_progress()

    def _poll_progress(self) -> None:
        """
        Lee mensajes de progreso y actualiza el status mientras el worker esté vivo.

        Importante: el chequeo de `worker.is_alive()` se hace al INICIO, no al final.
        Si el worker ya terminó, salimos sin tocar el status — así el mensaje
        final puesto por _on_fetch_done (vía _set_busy) no se sobrescribe con un
        "Procesando…" viejo que haya quedado en la cola.
        """
        if not (self.worker and self.worker.is_alive()):
            return
        try:
            while True:
                msg = self.progress_queue.get_nowait()
                self.status.configure(text=msg, foreground="#1f6feb")
        except Exception:
            pass
        # Reagendar. Si entre el check del inicio y ahora el worker murió,
        # la próxima iteración saldrá por el guard de arriba.
        self.root.after(120, self._poll_progress)

    def _fetch_worker(self, params: dict, on_done: Callable[[list[dict], str], None]) -> None:
        kind = params["kind"]

        def progress(msg: str, _count: int) -> None:
            self.progress_queue.put(msg)

        try:
            if kind == "today":
                progress("Consultando día actual…", 0)
                rows = fetch_day(date.today().isoformat())
            elif kind == "month":
                progress(
                    f"Consultando {calendar.month_name[params['month']]} {params['year']}…", 0
                )
                rows = fetch_month(params["year"], params["month"])
            else:  # range
                rows = fetch_range(params["start"], params["end"], on_progress=progress)
        except Exception as exc:  # noqa: BLE001
            on_done([], f"Error inesperado: {exc}")
            return
        if not rows:
            on_done([], "No se encontraron datos para el período seleccionado.")
        else:
            on_done(rows, f"Se obtuvieron {len(rows)} registro(s).")

    def _on_fetch_done(self, rows: list[dict], message: str) -> None:
        # Vuelve al hilo principal
        def _update() -> None:
            self.data = rows
            self._populate_table(rows)
            self._set_busy(False, message)
            self.count_label.configure(text=f"{len(rows)} fila(s)")
            # Drenar cualquier mensaje de progreso pendiente para que
            # una iteración tardía de _poll_progress no pinte nada viejo.
            try:
                while True:
                    self.progress_queue.get_nowait()
            except Exception:
                pass

        self.root.after(0, _update)

    # -- tabla ----------------------------------------------------------------
    def _clear_table(self) -> None:
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.count_label.configure(text="")

    def _populate_table(self, rows: list[dict]) -> None:
        self._clear_table()
        for i, row in enumerate(rows):
            tag = "even" if i % 2 == 0 else "odd"
            self.tree.insert(
                "",
                tk.END,
                values=(
                    self._format_fecha_display(row["fecha"]),
                    f"{row['compra']:.4f}",
                    f"{row['venta']:.4f}",
                ),
                tags=(tag,),
            )

    @staticmethod
    def _format_fecha_display(iso_str: str) -> str:
        try:
            d = datetime.strptime(iso_str[:10], "%Y-%m-%d").date()
            return d.strftime("%d/%m/%Y")
        except ValueError:
            return iso_str

    # -- portapapeles ---------------------------------------------------------
    def _get_rows_for_copy(self, selected_only: bool) -> list[list[str]]:
        if not self.data:
            return []
        if selected_only:
            sel = self.tree.selection()
            if not sel:
                return []
            indices = [self.tree.index(iid) for iid in sel]
            rows = [self.data[i] for i in indices]
        else:
            rows = self.data
        out: list[list[str]] = []
        out.append(["Fecha", "Compra", "Venta"])
        for r in rows:
            out.append(
                [
                    self._format_fecha_display(r["fecha"]),
                    f"{r['compra']:.4f}",
                    f"{r['venta']:.4f}",
                ]
            )
        return out

    def _copy_to_clipboard(self, rows: list[list[str]]) -> None:
        if not rows:
            messagebox.showinfo("Sin datos", "No hay filas para copiar.")
            return
        text = "\n".join("\t".join(r) for r in rows)
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        self.root.update()  # asegurar que el portapapeles persista
        n = len(rows) - 1  # descontar cabecera
        self.status.configure(
            text=f"{n} fila(s) copiada(s) al portapapeles.", foreground="#2da44e"
        )

    def _copy_selection(self) -> None:
        rows = self._get_rows_for_copy(selected_only=True)
        self._copy_to_clipboard(rows)

    def _copy_all(self) -> None:
        rows = self._get_rows_for_copy(selected_only=False)
        self._copy_to_clipboard(rows)

    # -- exportación Excel ----------------------------------------------------
    def _export_excel(self) -> None:
        if not self.data:
            messagebox.showinfo(
                "Sin datos", "No hay datos para exportar. Realiza primero una consulta."
            )
            return

        default_name = f"tipo_cambio_sunat_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filename = filedialog.asksaveasfilename(
            title="Guardar archivo Excel",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel (*.xlsx)", "*.xlsx"), ("Todos los archivos", "*.*")],
        )
        if not filename:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Tipo de Cambio SUNAT"

            # Estilos
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="1F4E78")
            header_align = Alignment(horizontal="center", vertical="center")
            cell_font = Font(name="Calibri", size=11)
            center = Alignment(horizontal="center", vertical="center")
            right = Alignment(horizontal="right", vertical="center")
            thin = Side(border_style="thin", color="BFBFBF")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            # Cabecera
            headers = ["Fecha", "Compra (S/)", "Venta (S/)"]
            ws.append(headers)
            for col_idx, _ in enumerate(headers, start=1):
                c = ws.cell(row=1, column=col_idx)
                c.font = header_font
                c.fill = header_fill
                c.alignment = header_align
                c.border = border

            # Filas
            for r in self.data:
                fecha_iso = r["fecha"][:10]
                try:
                    fecha_obj = datetime.strptime(fecha_iso, "%Y-%m-%d").date()
                except ValueError:
                    fecha_obj = fecha_iso
                ws.append([fecha_obj, float(r["compra"]), float(r["venta"])])

            # Formato de columnas
            last_row = ws.max_row
            for row_idx in range(2, last_row + 1):
                ws.cell(row=row_idx, column=1).number_format = "DD/MM/YYYY"
                ws.cell(row=row_idx, column=1).alignment = center
                ws.cell(row=row_idx, column=1).font = cell_font
                ws.cell(row=row_idx, column=1).border = border

                for col_idx in (2, 3):
                    c = ws.cell(row=row_idx, column=col_idx)
                    c.number_format = "0.0000"
                    c.alignment = right
                    c.font = cell_font
                    c.border = border

            # Anchos
            widths = (16, 16, 16)
            for col_idx, w in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(col_idx)].width = w

            # Congelar la cabecera
            ws.freeze_panes = "A2"

            wb.save(filename)
        except Exception as exc:  # noqa: BLE001
            messagebox.showerror("Error al exportar", f"No se pudo guardar el archivo:\n{exc}")
            return

        self.status.configure(text=f"Exportado a {filename}", foreground="#2da44e")
        if messagebox.askyesno("Exportación completa", f"Archivo guardado en:\n{filename}\n\n¿Abrir la carpeta?"):
            try:
                import os, subprocess
                folder = os.path.dirname(filename) or "."
                if hasattr(os, "startfile"):
                    os.startfile(folder)  # type: ignore[attr-defined]
                else:
                    subprocess.Popen(["xdg-open", folder])
            except Exception:  # noqa: BLE001
                pass


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main() -> None:
    """
    Si recibe argumentos de línea de comandos, despacha al modo CLI.
    Si no, abre la GUI.
    """
    if len(sys.argv) > 1:
        from cli import run
        sys.exit(run(sys.argv[1:]))
    root = tk.Tk()
    SunatTCApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
