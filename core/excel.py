"""Shared XLSX export for both the CLI (`cli.run --xlsx`) and the GUI
(`app.SunatTCApp._export_excel`).

Single source of truth for workbook styling (header, borders, number
formats, column widths, frozen header row). `openpyxl` is imported lazily
inside `write_xlsx`, so importing this module — or `core` itself — never
pays the openpyxl cost (app-startup-and-imports R1, S11).
"""
from __future__ import annotations

from datetime import datetime
from os import PathLike


def write_xlsx(rows: list[dict], path: str | PathLike[str]) -> None:
    """Escribe `rows` (dicts con fecha/compra/venta) a un XLSX con formato.

    Una `fecha` malformada cae al string crudo en vez de romper la
    exportación (guard defensivo heredado de cli._save_xlsx /
    app._export_excel). Lanza la excepción subyacente (OSError, ValueError,
    TypeError) si el archivo no puede escribirse; el caller la mapea a su
    contrato (CLI: exit 5; GUI: messagebox de error).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Tipo de Cambio SUNAT"

    # Estilos
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_align = Alignment(horizontal="center", vertical="center")
    cell_font = Font(name="Calibri", size=11)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    thin = Side(border_style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Cabecera
    headers = ["Fecha", "Compra (S/)", "Venta (S/)"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border

    # Filas
    for r in rows:
        fecha_iso = str(r["fecha"])[:10]
        try:
            fecha_obj = datetime.strptime(fecha_iso, "%Y-%m-%d").date()
        except ValueError:
            fecha_obj = fecha_iso
        ws.append([fecha_obj, float(r["compra"]), float(r["venta"])])

    # Formato de columnas
    last_row = ws.max_row
    for row_idx in range(2, last_row + 1):
        c1 = ws.cell(row=row_idx, column=1)
        c1.number_format = "DD/MM/YYYY"
        c1.alignment = center
        c1.font = cell_font
        c1.border = border
        for col_idx in (2, 3):
            c = ws.cell(row=row_idx, column=col_idx)
            c.number_format = "0.0000"
            c.alignment = right
            c.font = cell_font
            c.border = border

    # Anchos
    widths = (16, 16, 16)
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # Congelar la cabecera
    ws.freeze_panes = "A2"

    wb.save(path)
