"""
Modo CLI para sunat-tc-app.

Permite consultar el tipo de cambio de SUNAT desde la línea de comandos,
sin abrir la ventana. Útil para:
  - Tareas programadas (cron / Task Scheduler)
  - Pipelines que necesitan el TC en formato JSON/CSV
  - Uso remoto por SSH

Ejemplos:
    python app.py --today
    python app.py --yesterday
    python app.py --range 2024-01-01:2024-12-31 --json > tc_2024.json
    python app.py --month 2024-12 --xlsx tc_diciembre.xlsx
    python app.py --range 2024-01-01:2024-12-31 --csv tc_2024.csv
    python app.py --cache-info
    python app.py --clear-cache
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date, datetime, timedelta

# `openpyxl` is imported lazily inside `_save_xlsx` to keep `cli` import
# cheap and to avoid the GUI/CSV-only paths paying the openpyxl cost
# (app-startup-and-imports R1, S11).
from core import (
    CACHE_DB,
    FetchError,
    __version__,
    cache_clear,
    cache_count,
    cache_date_range,
    fetch_day,
    fetch_month,
    fetch_range,
    format_fetch_error,
    validate_month,
    validate_range,
)


def _build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="sunat-tc",
        description="Consulta el tipo de cambio de SUNAT desde la terminal.",
    )

    # Q8 (quality-lint-and-hardening D4): `--version` prints
    # `%(prog)s {__version__}` to stdout and exits 0 via argparse's
    # `action="version"` — fires during `parse_args`, so no query mode is
    # required and the fetch path is never reached.
    p.add_argument(
        "--version",
        action="version",
        version=f"%(prog)s {__version__}",
    )

    # Query mode (exactly one)
    query = p.add_mutually_exclusive_group()
    query.add_argument(
        "--today", action="store_true", help="TC del día actual"
    )
    query.add_argument(
        "--yesterday", action="store_true", help="TC del día anterior"
    )
    query.add_argument(
        "--range",
        metavar="F1:F2",
        help="Rango de fechas: F1=YYYY-MM-DD:F2=YYYY-MM-DD",
    )
    query.add_argument(
        "--month", metavar="YYYY-MM", help="Mes completo en formato YYYY-MM"
    )

    # Output options
    out = p.add_argument_group("salida")
    out.add_argument(
        "--json", action="store_true", help="Imprimir resultado en JSON a stdout"
    )
    out.add_argument(
        "--csv", metavar="ARCHIVO", help="Guardar resultado como CSV en ARCHIVO"
    )
    out.add_argument(
        "--xlsx", metavar="ARCHIVO", help="Guardar resultado como Excel en ARCHIVO"
    )

    # Cache
    cache = p.add_argument_group("caché")
    cache.add_argument(
        "--no-cache",
        action="store_true",
        help="Ignorar el caché local y consultar siempre a la API",
    )
    cache.add_argument(
        "--cache-info",
        action="store_true",
        help="Mostrar información del caché y salir",
    )
    cache.add_argument(
        "--clear-cache",
        action="store_true",
        help="Borrar el caché local y salir",
    )

    return p


def _print_table(rows: list[dict]) -> None:
    if not rows:
        print("(sin datos)", file=sys.stderr)
        return
    print(f"{'Fecha':<12} {'Compra':>10} {'Venta':>10}")
    print("-" * 34)
    for r in rows:
        fecha = r["fecha"]
        print(f"{fecha:<12} {r['compra']:>10.4f} {r['venta']:>10.4f}")


def _print_json(rows: list[dict]) -> None:
    json.dump(rows, sys.stdout, indent=2, ensure_ascii=False)
    sys.stdout.write("\n")


def _save_csv(rows: list[dict], path: str) -> None:
    # `utf-8-sig` writes the UTF-8 BOM (EF BB BF) so Microsoft Excel on
    # Windows opens the file with the correct encoding by default
    # (csv-export R1, S7).
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Fecha", "Compra", "Venta"])
        for r in rows:
            w.writerow([r["fecha"], f"{r['compra']:.4f}", f"{r['venta']:.4f}"])


def _save_xlsx(rows: list[dict], path: str) -> None:
    # Lazy import: keeps the cli module cheap for JSON/CSV/--help paths
    # and the import surface free of openpyxl (app-startup-and-imports R1, S11).
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Tipo de Cambio SUNAT"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_align = Alignment(horizontal="center", vertical="center")
    cell_font = Font(name="Calibri", size=11)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    thin = Side(border_style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["Fecha", "Compra (S/)", "Venta (S/)"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col_idx)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = border

    for r in rows:
        # Defend against a malformed `fecha` (theoretically possible after
        # manual DB edits or a future API change). Mirror `app.py:916-919`:
        # fall back to the raw string instead of crashing the export.
        # The outer `try/except (OSError, ValueError, TypeError)` below
        # converts any remaining error into the documented exit code 5.
        try:
            fecha_obj = datetime.strptime(r["fecha"], "%Y-%m-%d").date()
        except ValueError:
            fecha_obj = r["fecha"]
        ws.append([fecha_obj, float(r["compra"]), float(r["venta"])])

    last_row = ws.max_row
    for row_idx in range(2, last_row + 1):
        c1 = ws.cell(row=row_idx, column=1)
        c1.number_format = "DD/MM/YYYY"
        c1.alignment = center
        c1.font = cell_font
        c1.border = border
        for col_idx in (2, 3):
            c = ws.cell(row=row_idx, column=col_idx)
            c.number_format = "0.0000"
            c.alignment = right
            c.font = cell_font
            c.border = border

    widths = (16, 16, 16)
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    ws.freeze_panes = "A2"
    wb.save(path)


# Mapping from FetchError.kind to a user-facing Spanish message moved to
# `core/messages.py` (Q4 / fix-data-layer-bugs B3). Both the CLI and the
# GUI import `format_fetch_error` from `core` for a single source of truth.


def run(argv: list[str] | None = None) -> int:
    """Punto de entrada. Devuelve un código de salida."""
    parser = _build_parser()
    args = parser.parse_args(argv)

    # --- comandos que no requieren query ----------------------------------
    if args.clear_cache:
        n = cache_clear()
        print(f"Caché borrado. {n} entradas eliminadas.")
        return 0

    if args.cache_info:
        n = cache_count()
        dmin, dmax = cache_date_range()
        print(f"Entradas: {n}")
        print(f"Rango:   {dmin or '-'}  →  {dmax or '-'}")
        print(f"Archivo: {CACHE_DB}")
        return 0

    # --- algo de query ---------------------------------------------------
    if not (args.today or args.yesterday or args.range or args.month):
        parser.print_help(sys.stderr)
        return 1

    use_cache = not args.no_cache

    try:
        if args.today:
            rows = fetch_day(date.today().isoformat(), use_cache=use_cache)
        elif args.yesterday:
            d = date.today() - timedelta(days=1)
            rows = fetch_day(d.isoformat(), use_cache=use_cache)
        elif args.range:
            try:
                s, e = args.range.split(":")
                sd = datetime.strptime(s, "%Y-%m-%d").date()
                ed = datetime.strptime(e, "%Y-%m-%d").date()
            except ValueError:
                print(
                    "Error: --range debe tener formato YYYY-MM-DD:YYYY-MM-DD",
                    file=sys.stderr,
                )
                return 2
            v = validate_range(sd, ed)
            if isinstance(v, str):
                print(f"Error: {v}", file=sys.stderr)
                return 2
            sd, ed = v
            rows = fetch_range(sd, ed, use_cache=use_cache)
        else:  # args.month
            try:
                y_str, m_str = args.month.split("-")
                yr, mo = int(y_str), int(m_str)
            except ValueError:
                print("Error: --month debe tener formato YYYY-MM", file=sys.stderr)
                return 2
            v = validate_month(yr, mo)
            if isinstance(v, str):
                print(f"Error: {v}", file=sys.stderr)
                return 2
            yr, mo = v
            rows = fetch_month(yr, mo, use_cache=use_cache)
    except FetchError as exc:
        print(f"Error: {format_fetch_error(exc)}", file=sys.stderr)
        return 4

    if not rows:
        print("No se encontraron datos para el período solicitado.", file=sys.stderr)
        return 3

    # --- salida ----------------------------------------------------------
    try:
        if args.json:
            _print_json(rows)
        elif args.csv:
            _save_csv(rows, args.csv)
            print(f"Guardado: {args.csv}  ({len(rows)} filas)", file=sys.stderr)
        elif args.xlsx:
            _save_xlsx(rows, args.xlsx)
            print(f"Guardado: {args.xlsx}  ({len(rows)} filas)", file=sys.stderr)
        else:
            _print_table(rows)
            print(f"\n{len(rows)} fila(s).", file=sys.stderr)
    except (OSError, ValueError, TypeError) as exc:
        print(f"Error al escribir: {exc}", file=sys.stderr)
        return 5

    return 0


def main() -> None:
    sys.exit(run())


if __name__ == "__main__":
    main()
